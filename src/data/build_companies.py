"""Regenerates src/data/companies.json from src/data/company_wise_sheet.xlsx.

One sheet per company. Run from the repo root:  python3 src/data/build_companies.py
Requires: openpyxl.  See CompanyPrep.tsx — COMPANY_DOMAINS must stay in sync with the
company ids produced here.
"""
import json, os, re
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SRC = os.path.join(ROOT, 'src/data/company_wise_sheet.xlsx')
OUT = os.path.join(ROOT, 'src/data/companies.json')
DSA = os.path.join(ROOT, 'src/data/dsaSheets.json')

# slug -> (topic, difficulty, title) index built from the curated DSA sheet
slugidx = {}
_dsa = json.load(open(DSA))
_SLUG = re.compile(r'(?:leetcode\.com/problems/|geeksforgeeks\.org/problems/|geeksforgeeks\.org/dsa/|geeksforgeeks\.org/[a-z\-]+/|geeksforgeeks\.org/)([a-z0-9\-]+)')
for _s in _dsa['sheets']:
    for _lv in _s['levels']:
        for _t in _lv['topics']:
            for _p in _t['problems']:
                for _url in (_p.get('platforms') or {}).values():
                    _m = _SLUG.search((_url or '').lower())
                    if _m:
                        slugidx.setdefault(_m.group(1), (_t['name'], _p['difficulty'], _p['title']))

# sheet name -> (display name, id, domain or None)
COMPANIES = {
    'Accenture': ('Accenture', 'accenture', 'accenture.com'),
    'Accolite': ('Accolite Digital', 'accolite', 'accolite.com'),
    'Alstom': ('Alstom', 'alstom', 'alstom.com'),
    'Amazon': ('Amazon', 'amazon', 'amazon.com'),
    'Amex': ('American Express', 'american-express', 'americanexpress.com'),
    'Appian': ('Appian', 'appian', 'appian.com'),
    'Apple': ('Apple', 'apple', 'apple.com'),
    'Aptiv': ('Aptiv', 'aptiv', 'aptiv.com'),
    'Arcesium': ('Arcesium', 'arcesium', 'arcesium.com'),
    'Arista': ('Arista Networks', 'arista-networks', 'arista.com'),
    'Aspire sys': ('Aspire Systems', 'aspire-systems', 'aspiresys.com'),
    'Athena health': ('Athenahealth', 'athena-health', 'athenahealth.com'),
    'Bank of america': ('Bank of America', 'bank-of-america', 'bankofamerica.com'),
    'barclays': ('Barclays', 'barclays', 'barclays.com'),
    'baton': ('Baton Systems', 'baton-systems', 'batonsystems.com'),
    'bny': ('BNY', 'bny', 'bny.com'),
    'checktronix': ('Checktronix', 'checktronix', 'checktronix.com'),
    'chronus': ('Chronus', 'chronus', 'chronus.com'),
    'DE shaw': ('D. E. Shaw & Co.', 'de-shaw', 'deshaw.com'),
    'citi': ('Citi', 'citi', 'citi.com'),
    'deloitte': ('Deloitte', 'deloitte', 'deloitte.com'),
    'dover': ('Dover', 'dover', 'dovercorporation.com'),
    'enphase': ('Enphase Energy', 'enphase', 'enphase.com'),
    'eucloid data sol': ('Eucloid Data Solutions', 'eucloid', 'eucloid.com'),
    'fidelity': ('Fidelity Investments', 'fidelity', 'fidelity.com'),
    'gen': ('Gen Digital', 'gen-digital', 'gendigital.com'),
    'global analytics': ('Global Analytics', 'global-analytics', 'globalanalytics.com'),
    'goldman sachs': ('Goldman Sachs', 'goldman-sachs', 'goldmansachs.com'),
    'heftin ai': ('Heftin AI', 'heftin-ai', None),
    'HSBC': ('HSBC', 'hsbc', 'hsbc.com'),
    'Hubstream': ('Hubstream', 'hubstream', 'hubstream.com'),
    'IDFC': ('IDFC First Bank', 'idfc-first-bank', 'idfcfirstbank.com'),
    'Infibeam avenues': ('Infibeam Avenues', 'infibeam-avenues', 'infibeamav.com'),
    'infinera': ('Infinera', 'infinera', 'infinera.com'),
    'KLA': ('KLA', 'kla', 'kla.com'),
    'Khoros': ('Khoros', 'khoros', 'khoros.com'),
    'Loadshare ntk': ('Loadshare Networks', 'loadshare-networks', 'loadshare.net'),
    'LTI': ('LTIMindtree', 'ltimindtree', 'ltimindtree.com'),
    'Mathworks': ('MathWorks', 'mathworks', 'mathworks.com'),
    'Micron tech': ('Micron Technology', 'micron', 'micron.com'),
    'microsoft': ('Microsoft', 'microsoft', 'microsoft.com'),
    'morgan stanley': ('Morgan Stanley', 'morgan-stanley', 'morganstanley.com'),
    'natwest': ('NatWest Group', 'natwest', 'natwestgroup.com'),
    'NCR': ('NCR Voyix', 'ncr-voyix', 'ncrvoyix.com'),
    'nokia': ('Nokia', 'nokia', 'nokia.com'),
    'NVDIA': ('NVIDIA', 'nvidia', 'nvidia.com'),
    'Optum': ('Optum', 'optum', 'optum.com'),
    'oracle': ('Oracle', 'oracle', 'oracle.com'),
    'Q2': ('Q2 Software', 'q2', 'q2.com'),
    'Quantiphi': ('Quantiphi', 'quantiphi', 'quantiphi.com'),
    'Ramco': ('Ramco Systems', 'ramco-systems', 'ramco.com'),
    'RtBrick': ('RtBrick', 'rtbrick', 'rtbrick.com'),
    'Samsung R&D': ('Samsung R&D Institute', 'samsung-rd', 'samsung.com'),
    'SAP Labs': ('SAP Labs', 'sap-labs', 'sap.com'),
    'Software AG': ('Software AG', 'software-ag', 'softwareag.com'),
    'Tekion': ('Tekion', 'tekion', 'tekion.com'),
    'Trimble': ('Trimble', 'trimble', 'trimble.com'),
    'Vegrow': ('Vegrow', 'vegrow', 'vegrow.in'),
    'Verizon': ('Verizon', 'verizon', 'verizon.com'),
    'Versa ntk': ('Versa Networks', 'versa-networks', 'versa-networks.com'),
    'Vivriti cap': ('Vivriti Capital', 'vivriti-capital', 'vivriticapital.com'),
    'VISA': ('Visa', 'visa', 'visa.com'),
    'Walmart': ('Walmart', 'walmart', 'walmart.com'),
    'Wells Fargo': ('Wells Fargo', 'wells-fargo', 'wellsfargo.com'),
    'WD': ('Western Digital', 'western-digital', 'westerndigital.com'),
    'Wex fintech': ('WEX', 'wex', 'wexinc.com'),
    'wipro': ('Wipro', 'wipro', 'wipro.com'),
    'Zoho': ('Zoho', 'zoho', 'zoho.com'),
}

JUNK = {'none', 'nothing', 'nil', 'na', 'n/a', '-', '', 'null'}

def clean(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None

def norm_year(v):
    if v is None:
        return ''
    if isinstance(v, (int, float)):
        return str(int(v))
    s = str(v).strip()
    yrs = re.findall(r'(20\d{2})', s)
    if yrs:
        return min(yrs)
    m = re.search(r'(20\d{2})\s*/\s*(\d{2})', s)
    if m:
        return m.group(1)
    return re.sub(r'\.0$', '', s)

def norm_diff(v):
    if not v:
        return None
    s = str(v).strip().lower()
    if 'hard' in s:
        return 'Hard'
    if 'medium' in s or 'moderate' in s:
        return 'Medium'
    if 'easy' in s:
        return 'Easy'
    return None

def norm_type(round_txt):
    r = (round_txt or '').strip().lower()
    if r == 'oa' or 'online' in r or 'coding round' in r or 'coding assessment' in r \
       or 'coding test' in r or 'written' in r or 'aptitude' in r or 'offline coding' in r \
       or 'programming' in r or 'hackerrank' in r:
        return 'OA'
    return 'Interview'

SLUG_RE = re.compile(r'(?:leetcode\.com/problems/|geeksforgeeks\.org/problems/|geeksforgeeks\.org/dsa/|geeksforgeeks\.org/[a-z\-]+/|geeksforgeeks\.org/)([a-z0-9\-]+)')

def link_slug(u):
    if not u:
        return None
    m = SLUG_RE.search(u.lower())
    return m.group(1) if m else None

# ordered (regex, topic)
TOPIC_RULES = [
    (r'\bsql\b|\bdbms\b|\brdbms\b|database|normal ?form|normali[sz]|\bjoin\b|subquery|indexing in|query processing|truncate|foreign key|primary key|referential integ|\bddl\b|\bdml\b|\bacid\b|\bbase properties|two[\s\-]phase locking|functional dependency|\btrigger\b|group by|customers who have not|data warehous|one[\s\-]to[\s\-]many', 'Databases'),
    (r'\boop\b|object[\s\-]oriented|polymorph|inherit|encapsul|abstraction|method (overload|overrid)|function (overload|overrid)|access modifier|namespace|singleton|solid (design )?principle|design pattern|\bhttp\b|\bftp\b|sftp|\bsmtp\b|\btcp\b|\budp\b|\bosi\b|\bdns\b|operating system|\bthread|concurrency|parallelism|context switch|virtual memory|\bpaging\b|segmentation|deadlock|cpu schedul|memory layout|garbage collect|wrapper class|pass by (value|reference)|exception|\bstl\b|standard template|template library|malloc|calloc|deep copy|shallow copy|\bpointer|semiconductor|mosfet|\bdiode\b|pn junction|ohm|forward bias|reverse bias|multimeter|power supply|\bbios\b|risc|cisc|endian|32[\s\-]bit|64[\s\-]bit|platform independ|servlet|virtual dom|react|node\.?js|promise|async|event loop|middleware|mongodb|firebase|rest api|\bjwt\b|docker|kubernetes|ci/cd|jenkins|ansible|\bgit\b|github|saas|paas|iaas|scaling|gradient descent|linear regression|machine learning|deep learning|overfitting|underfitting|supervised|unsupervised|classification and regression|\blime\b|\bshap\b|explainability|sessions and cookies|truthy|falsy|frontend and backend|\bsdlc\b|software development life|time complexity|complexity of|\bvoid\b and \bnull\b|stdio\.h|memory of|\bcookies\b|avl tree|red[\s\-]black tree|b[\s\-]tree', 'CS Fundamentals'),
    (r'design (a|an|concurrent|scalable|the)|[\w]+[\s\-]like (booking|application|system)|booking (application|system)|ticket booking|inventory management|weather app|chess game|uber |redbus|bookmyshow|redbus|handle (payment|too many|concurrent|simultaneous)|rate limit|handle heavy load|api to handle', 'System Design'),
    (r'linked[\s\-]?list|\bdll\b|reverse (a )?list|reverse nodes|flatten(ing)? a? ?(multilevel|linked)|nth node from|node from (the )?end|loop in a? ?linked', 'Linked Lists'),
    (r'binary search tree|\bbst\b|insert into a binary search|search in a binary search|delete node in a bst', 'Binary Search Trees'),
    (r'binary tree|tree traversal|root[\s\-]to[\s\-]leaf|subtree|lowest common ancestor|\blca\b|level[\s\-]order|zigzag|inorder|preorder|postorder|tree height|tree diameter|mirror.*tree|tree.*mirror|symmetric tree|invert binary|duplicate subtree|maximum sum bst|boundary traversal|special nodes in tree|alternate nodes', 'Binary Tree'),
    (r'\btrie\b|autocomplete|prefix search|search suggestion|26[\s\-]ary', 'Trie'),
    (r'\bgraph\b|island|connected component|topolog|dijkstra|\bbfs\b|\bdfs\b|course schedule|critical connection|adjacency|rotten? orange|rotting orange|detonate|bombs|enclosed land|number of islands|hamiltonian|seven bridges|konigsberg|minimum (pipe|spanning)|spanning tree|closest pair of points|travel(l)?ing salesman|\btsp\b|detect cycle in .*graph|count zero request|server assignment', 'Graphs'),
    (r'dynamic programming|\bdp\b|knapsack|coin change|edit distance|longest common|longest increasing|\blis\b|palindrom.*partition|jump game|unique paths|maximum sum rectangle|weighted (interval|schedul)|matrix path|max reward|egg drop|climb|nth stair|count ways to reach|ways to reach|min(imum)? (insertion|deletion) steps|house[\s\-]?robb|maximum length chain|chain of pairs|weighted schedul|nth term of', 'Dynamic Programming'),
    (r'greedy|activity selection|interval schedul|merge intervals|overlapping interval|minimum (number of )?platform|meeting rooms?|meetings? in one room|minimum meeting|gas station|candy', 'Greedy Algorithms'),
    (r'\bheap\b|priority queue|kth largest|kth smallest|top k|k workers|\bmedian\b|hire k|least[\s\-]loaded|smallest range covering', 'Heap'),
    (r'\bstack\b|infix|postfix|prefix expression|valid parenthes|balanced paren|match paren|next greater|histogram|largest rectangle|simplify path|push and pop|two stacks', 'Stacks'),
    (r'\bqueue\b|\bdeque\b|circular queue|implement queue|time needed to buy|buy tickets', 'Queues'),
    (r'backtrack|permutation|combination|\brecursi|josephus|n[\s\-]?queens|sudoku|generate all|generate binary string|subsets|rat (in|and)|knight move|tower of hanoi|all binary strings|word search|maze|nearest exit', 'Recursion & Backtracking'),
    (r'sliding window|substring without repeat|longest substring|subarray|three consecutive|consecutive transaction|consecutive values|consecutive differences|k distinct|even[\s\-]length substring', 'Sliding Window'),
    (r'prefix sum|nice subarrays|subarrays? with (given )?sum', 'Hashing'),
    (r'\bhash|anagram|frequenc|duplicate|first repeat|first non[\s\-]?repeat|two occurrences|majority element|\bdistinct\b|appearing only once|appears? only once|intersection of two arrays|single (non[\s\-]duplicate|number)|repeated (word|element)|count elements|most frequent', 'Hashing'),
    (r'two pointer|\b3\s?sum\b|two sum|trapping rain|container with most|\bpair(s)? (with|summing)|triplet|closest pair between|merge (two )?sorted (array|list)|remove duplicates from sorted|move zero|sort colors|dutch national', 'Two Pointers'),
    (r'bitwise|bit manipulation|\bxor\b|power of two|set bits?|count bits|binary representation|gray code|swap.*without (temp|conditional|using)|without (a )?temp(orary)? variable|division and multiplication without|without (direct )?arithmetic|divide two integer|without operator|add(ition)? (of )?two integers.*bit|missing coordinate|missing vertex|xor property', 'Bit Manipulation'),
    (r'puzzle|torch and bridge|balance scale|heavier ball|weight.*ball|\briddle|water jug|water (dispenser|cups)|fill (the )?cups|bill split|lift direction|elevator floor|\blamp[\s\-]?post|lamp states', 'Puzzles'),
    (r'binary search|aggressive cows|koko eating|find position.*sorted|search insert|infinite array|allocate|painter|book allocation|first and last (occurrence|digit)|rotated sorted|minimize the (max|largest)|split .* minimize|minimum .* to (fill|reach)|find the winner|removing (at most )?\d? ?stones|stone (removal|game)', 'Searching'),
    (r'\bsort\b|sorting|bubble sort|merge sort|quick sort|selection sort|insertion sort|0s.*1s.*2s|sort characters|sort .* by (frequency|squared|their)|roman numeral|sorted position', 'Sorting'),
    (r'\bstring\b|substring|palindrom|reverse (words|a string|the string|integer|sentence)|reverse.*sentence|\bcharacter|lexicograph|\bip\b|ipv4|valid(ate)? ip|wildcard match|word break|string generation|run[\s\-]length|text justification|justif|replace .* question mark|latest valid date|two[\s\-]digit substring|largest .* substring|backspace|rearrange.*string|shift(ing)? character|remove k digits|remove all adjacent|candy crush|consecutive character removal', 'Strings'),
    (r'\bprime\b|fibonacci|\bgcd\b|\blcm\b|\bdigit|geometr|combinator|perfect square|number to (words|english)|to english words|modulo|divisib|triangle|parallelogram|coordinate|\bsquare\b|factor|happy (array|number)|weird number|pascal|sum of digits|first n primes|reachable by a bishop|chessboard cells|\bmath\b|count numbers|occurrences of|convert number|binary .* number|number spiral|diagonal|calendar|years, months|lunar phase|date', 'Math'),
    (r'matrix|grid|spiral|\brotate\b|kadane|maximum subarray|max(imum)? consecutive|missing (number|element|and repeated|point|positive)|first missing positive|second[\s\-](largest|most)|largest (two |three )?element|min(imum)? and max(imum)?|maximum.*array|rearrange|partition negatives|\barray\b|max grayness|run length|ordered stream|tic[\s\-]?tac[\s\-]?toe|max(imum)? (of )?two|max\(a|prefix|suffix|alternate index|sum of (array )?element|filter array|distribute elements|sum both diagonal|nearest (city|free)|parking|overtakes|max(imum)? (absolute )?difference', 'Arrays'),
    (r'\bexplain\b|\bcompare\b|difference between|what (is|are|happens)|how (does|do|to)|write (a |example )?(sql|code|query)|real[\s\-]world', 'CS Fundamentals'),
]
COMPILED = [(re.compile(p), t) for p, t in TOPIC_RULES]

def classify_topic(title, pattern, slug):
    if slug and slug in slugidx:
        return slugidx[slug][0]
    hay = ' '.join(x for x in [pattern, title] if x).lower()
    for rx, topic in COMPILED:
        if rx.search(hay):
            return topic
    return 'Miscellaneous'

def prettify_slug(slug):
    slug = re.sub(r'\d{3,}$', '', slug).rstrip('-')
    words = slug.replace('-', ' ').split()
    small = {'a', 'an', 'the', 'of', 'in', 'to', 'and', 'or', 'with', 'from', 'via', 'by'}
    out = []
    for i, w in enumerate(words):
        if w in ('i', 'ii', 'iii', 'iv', 'v', 'vi'):
            out.append(w.upper())
        elif i > 0 and w in small:
            out.append(w)
        else:
            out.append(w.capitalize())
    return ' '.join(out)

def derive_title(link, slug):
    if slug and slug in slugidx:
        return slugidx[slug][2]
    if slug:
        return prettify_slug(slug)
    if link:
        m = re.search(r'problem[=/](\d+)', link)
        host = re.sub(r'^https?://(www\.)?', '', link).split('/')[0]
        if m:
            return f'{host} problem {m.group(1)}'
        return f'Problem ({host})'
    return 'Untitled problem'

wb = openpyxl.load_workbook(SRC, data_only=True)
result = []
unclassified = []
report = []

for sheet in wb.sheetnames:
    if sheet not in COMPANIES:
        report.append(f'SKIP (not in map): {sheet}')
        continue
    name, cid, domain = COMPANIES[sheet]
    ws = wb[sheet]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]

    # locate header
    header = None
    hidx = -1
    for i, r in enumerate(rows):
        cells = [clean(c) for c in r]
        low = [c.lower() if c else '' for c in cells]
        if 's no' in low or ('name' in low and ('link' in low or 'level' in low or 'pattern type' in low)):
            header, hidx = cells, i
            break
        if low[:4] == ['link', 'difficulty', 'round', 'year']:
            header, hidx = cells, i
            break
    linkonly = header is not None and (header[0] or '').lower() == 'link'

    if header is None:
        # positional fallback: S No, Name, Level, Link, Asked in, year, Pattern Type, Company
        colmap = {'name': 1, 'level': 2, 'link': 3, 'round': 4, 'year': 5, 'pattern': 6}
        data_rows = rows
    elif linkonly:
        colmap = {'link': 0, 'level': 1, 'round': 2, 'year': 3}
        data_rows = rows[hidx + 1:]
    else:
        colmap = {}
        for j, h in enumerate(header):
            if not h:
                continue
            hl = h.lower()
            if hl == 'name':
                colmap['name'] = j
            elif hl in ('level', 'difficulty'):
                colmap['level'] = j
            elif hl == 'link':
                colmap['link'] = j
            elif hl in ('asked in', 'round'):
                colmap['round'] = j
            elif hl in ('year',):
                colmap['year'] = j
            elif hl == 'pattern type':
                colmap['pattern'] = j
        data_rows = rows[hidx + 1:]

    def get(r, key):
        j = colmap.get(key)
        if j is None or j >= len(r):
            return None
        return clean(r[j])

    questions = []
    seen = set()
    for r in data_rows:
        if not any(c is not None and str(c).strip() for c in r):
            continue
        nm = get(r, 'name')
        lk = get(r, 'link')
        # junk marker rows
        joined = ' '.join(str(c).strip().lower() for c in r if c is not None and str(c).strip())
        if joined in JUNK:
            continue
        if lk and lk.strip() in ('-', '—'):
            lk = None
        has_link = bool(lk and lk.lower().startswith('http'))
        slug = link_slug(lk) if has_link else None
        if not nm and not has_link:
            continue
        if linkonly:
            nm = derive_title(lk, slug)
        if not nm:
            nm = derive_title(lk, slug)

        dkey = (re.sub(r'[^a-z0-9]', '', nm.lower()), link_slug(lk) or (lk or '').lower())
        if dkey in seen:
            continue
        seen.add(dkey)

        pattern = get(r, 'pattern')
        diff = norm_diff(get(r, 'level'))
        if not diff and slug and slug in slugidx:
            diff = slugidx[slug][1]
        if not diff:
            diff = 'Medium'
        topic = classify_topic(nm, pattern, slug)
        if topic == 'Miscellaneous':
            unclassified.append((sheet, nm, pattern))

        q = {
            'title': nm,
            'link': lk if has_link else '',
            'type': norm_type(get(r, 'round')),
            'year': norm_year(get(r, 'year')),
            'id': f'{cid}-q{len(questions) + 1}',
            'topic': topic,
            'difficulty': diff,
        }
        questions.append(q)

    if not questions:
        report.append(f'SKIP (no questions): {sheet}')
        continue

    report.append(f'{name:28s} {len(questions):3d} q  ({sheet})')
    result.append({
        'id': cid,
        'name': name,
        'questions': questions,
        'oaFormat': [],
    })

result.sort(key=lambda c: (-len(c['questions']), c['name']))

with open(OUT, 'w') as f:
    json.dump(result, f, indent=2, ensure_ascii=False)
    f.write('\n')

print('\n'.join(report))
print(f'\nTOTAL companies: {len(result)}, total questions: {sum(len(c["questions"]) for c in result)}')
print(f'\nUNCLASSIFIED ({len(unclassified)}):')
for s, n, p in unclassified:
    print(f'  [{s}] {n!r} pattern={p!r}')

# COMPANY_DOMAINS block for CompanyPrep.tsx (paste + keep sorted)
doms = {v[1]: v[2] for v in COMPANIES.values() if v[2]}
print('\nCOMPANY_DOMAINS:')
for k in sorted(doms):
    key = k if re.match(r'^[a-z][a-z0-9]*$', k) else f"'{k}'"
    print(f"  {key}: '{doms[k]}',")
